"""Report composer. Every claim carries a question-ID/answer (and evidence) provenance link.

Content is deterministic: nothing is inferred beyond what answers and rules state.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
from typing import Any

SEV = {"critical": 0, "high": 1, "medium": 2, "low": 3}
PRELIMINARY_BANNER = "PRELIMINARY - NOT REVIEWED BY CLOUDSEALS. Do not rely on this document."
DISCLAIMER = ("Findings are candidates produced by deterministic rules from customer-provided answers and are subject to "
              "human review. Regulatory references are screening aids and not a legal determination of applicability.")
SEVERITY_ORDER = ("critical", "high", "medium", "low")


_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\u202a-\u202e\u2066-\u2069\ud800-\udfff\ufffe\uffff]")


def clean(text: Any) -> str:
    """Strip control and bidi-override characters that corrupt DOCX XML or visually spoof content."""
    return _BAD.sub("", str(text))


def _domain(question_id: str) -> str:
    return question_id.split("-")[1]


def _claim(reg, qid: str, answers: dict, answer_meta: dict) -> dict:
    q = reg.questions[qid]
    meta = answer_meta.get(qid, {})
    return {"question_id": qid, "question": q["question"], "answer": answers.get(qid),
            "assurance": "Self-attested" if not meta.get("evidence_refs") else "Supported by document (pending review)",
            "evidence_refs": meta.get("evidence_refs", []), "comment": meta.get("comment", "")}


def build_model(reg, assessment: dict, answers: dict[str, Any], answer_meta: dict[str, dict], findings: list[dict],
                not_assessed: list[dict], status: str, review: dict | None = None) -> dict:
    """status 'published' => final; anything else => visibly PRELIMINARY."""
    final = status == "published"
    by_domain: dict[str, list[str]] = {}
    for qid in answers:
        if qid in reg.questions:
            by_domain.setdefault(_domain(qid), []).append(qid)
    fs = sorted(findings, key=lambda f: (SEV.get(f.get("severity_hint", "low"), 9), f["rule_id"]))
    sections = []
    for spec in reg.report_sections:
        doms = spec["source_domains"]
        items: list[dict] = []
        notes: list[str] = []
        if doms == ["ALL"]:
            pass
        else:
            for d in doms:
                items += [_claim(reg, q, answers, answer_meta) for q in sorted(by_domain.get(d, []))]
        sec_findings = fs if doms == ["ALL"] else [f for f in fs if {c.split("-")[1] for c in f["control_ids"]} & set(doms)]
        sid = spec["id"]
        if sid == "REP-01":
            sev = {s: sum(f.get("severity_hint") == s for f in fs) for s in SEVERITY_ORDER}
            notes.append(f"{len(fs)} candidate finding(s): " + ", ".join(f"{n} {s}" for s, n in sev.items() if n) if fs
                         else "No candidate findings were raised from the submitted answers.")
            sec_findings = fs
        if sid == "REP-08" or sid == "REP-09":
            notes.append("Candidate screening only; legal review required before relying on applicability.")
            hits = [x for x in reg.crosswalk if any(r in answers for r in x["question_refs"])]
            items += [{"reference_id": x["reference_id"], "subject": f"{x['article']} - {x['subject']}",
                       "applicability": x["applicability"] + " (candidate)", "scope_note": x["scope_note"],
                       "answers": {r: answers[r] for r in x["question_refs"] if r in answers},
                       "source_url": x["source_url"]} for x in hits]
        if sid == "REP-11":
            notes.append("Roadmap ordered by severity hint; each action traces to a finding rule.")
            items = [{"rule_id": f["rule_id"], "action": f["recommendation"], "severity": f["severity_hint"],
                      "controls": f["control_ids"]} for f in fs]
            sec_findings = []
        if sid == "REP-15":
            items = [{"questionnaire_version": assessment["questionnaire_version"], "registry_sha256": assessment["registry_sha256"],
                      "registry_status": reg.raw["status"],
                      "rule_versions": {r["rule_id"]: r["rule_version"] for r in reg.risk_rules}}]
        if sid == "REP-16":
            items = [{"rule_id": n["rule_id"], "title": n["title"], "missing": n["missing"], "unknown": n["unknown"]} for n in not_assessed]
            items += [{"question_id": q, "question": reg.questions[q]["question"], "answer": "Unknown"}
                      for q, v in sorted(answers.items()) if v == "Unknown" and q in reg.questions]
        sections.append({"id": sid, "title": spec["title"], "purpose": spec["content"], "notes": notes,
                         "items": items, "findings": sec_findings})
    model = {"title": f"AI Architecture Diagnosis - {assessment['ai_system_name']}",
             "assessment_id": assessment["assessment_id"], "status": "FINAL" if final else "PRELIMINARY",
             "banner": None if final else PRELIMINARY_BANNER, "disclaimer": DISCLAIMER,
             "review": review if final else None,
             "provenance": {"questionnaire_version": assessment["questionnaire_version"],
                            "registry_sha256": assessment["registry_sha256"],
                            "answer_manifest_digest": (assessment.get("snapshot") or {}).get("answer_manifest_digest")},
             "sections": sections}
    model["content_digest"] = hashlib.sha256(json.dumps(model, sort_keys=True, default=str).encode()).hexdigest()
    return model


def _flat(item: dict) -> str:
    return clean(_flat_raw(item))


def _flat_raw(item: dict) -> str:
    if "question_id" in item and "answer" in item and "assurance" in item:
        ev = f"; evidence: {', '.join(item['evidence_refs'])}" if item["evidence_refs"] else ""
        ans = ", ".join(map(str, item["answer"])) if isinstance(item["answer"], list) else item["answer"]
        return f"[{item['question_id']}] {item['question']} -> {ans} ({item['assurance']}{ev})"
    return "; ".join(f"{k}: {v}" for k, v in item.items())


def _finding_line(f: dict) -> str:
    cites = ", ".join(f"{c['question_id']}={c['answer']}" for c in f["citations"])
    return clean(f"{f['rule_id']} v{f['rule_version']} [{f['severity_hint']}, {f['status']}, confidence {f['confidence']}] "
            f"{f['title']} - {f['recommendation']} (basis: {cites})")


def render_docx(model: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    d = Document()
    d.core_properties.title = model["title"]
    d.add_heading(clean(model["title"]), 0)
    if model["banner"]:
        p = d.add_paragraph()
        r = p.add_run(model["banner"])
        r.bold = True
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    d.add_paragraph(f"Status: {model['status']}  |  Assessment: {model['assessment_id']}")
    d.add_paragraph(f"Questionnaire {model['provenance']['questionnaire_version']}  |  Registry SHA256 {model['provenance']['registry_sha256'][:16]}...")
    d.add_paragraph(model["disclaimer"]).runs[0].italic = True
    for s in model["sections"]:
        d.add_heading(f"{s['id']} {s['title']}", 1)
        d.add_paragraph(s["purpose"]).runs[0].italic = True
        for n in s["notes"]:
            d.add_paragraph(clean(n))
        for it in s["items"]:
            d.add_paragraph(_flat(it), style="List Bullet").runs[0].font.size = Pt(9)
        for f in s["findings"]:
            d.add_paragraph(_finding_line(f), style="List Number").runs[0].font.size = Pt(9)
        if not (s["items"] or s["findings"] or s["notes"]):
            d.add_paragraph("No answers recorded for this section.")
    d.add_paragraph(f"Content digest: {model['content_digest']}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_pdf(model: dict) -> bytes:
    from xml.sax.saxutils import escape as esc

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    ss = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=model["title"], author="CloudSeals TrustOps")
    st: list = [Paragraph(esc(clean(model["title"])), ss["Title"])]
    if model["banner"]:
        st.append(Paragraph(f"<font color='#C00000'><b>{esc(model['banner'])}</b></font>", ss["Normal"]))
    st += [Paragraph(esc(f"Status: {model['status']} | Assessment {model['assessment_id']}"), ss["Normal"]),
           Paragraph(f"<i>{esc(model['disclaimer'])}</i>", ss["Normal"]), Spacer(1, 12)]
    for s in model["sections"]:
        st += [Paragraph(esc(f"{s['id']} {s['title']}"), ss["Heading2"]), Paragraph(f"<i>{esc(s['purpose'])}</i>", ss["Normal"])]
        st += [Paragraph(esc(clean(n)), ss["Normal"]) for n in s["notes"]]
        st += [Paragraph("&bull; " + esc(_flat(i)), ss["BodyText"]) for i in s["items"]]
        st += [Paragraph(esc(_finding_line(f)), ss["BodyText"]) for f in s["findings"]]
        st.append(Spacer(1, 8))
    st.append(Paragraph(esc(f"Content digest: {model['content_digest']}"), ss["Code"]))
    doc.build(st)
    return buf.getvalue()
